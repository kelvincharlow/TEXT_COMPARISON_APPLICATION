"""Preview authorization, retryable cache generation and real page rendering."""
import shutil
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from docx import Document
from backend.app.previews import prepare_preview, render_page
from backend.tests.test_reviews import ReviewFixture


class PreviewTests(ReviewFixture):
    def fake_pdf(self, source, destination):
        destination.write_bytes(b'PDF placeholder')
        return 2 if source.name == 'original.docx' else 3

    def test_preview_access_matches_document_access_and_validates_pages(self):
        p = self.comparison()
        url = f"/api/v1/comparisons/{p['comparison_id']}/preview"
        with patch('backend.app.previews.render_pdf', side_effect=self.fake_pdf) as render:
            for client in (self.outsider, self.admin):
                self.assertEqual(client.post(url).status_code, 404)
                self.assertEqual(client.get(url + '/original/pages/1').status_code, 404)
            self.assertEqual(render.call_count, 0)
            response = self.reviewer.post(url)
            self.assertEqual(response.status_code, 200, response.text)
            self.assertEqual(response.json()['original']['page_count'], 2)
            self.assertEqual(response.json()['redline']['page_count'], 3)
            self.assertEqual(self.creator.post(url).status_code, 200)
            self.assertEqual(render.call_count, 2)
        for suffix in ('original/pages/0', 'original/pages/3', 'redline/pages/4'):
            self.assertEqual(self.reviewer.get(url + '/' + suffix).status_code, 404)
        self.assertEqual(self.reviewer.get(url + '/revised/pages/1').status_code, 422)
        def png(args):
            Path(args[-1] + '.png').write_bytes(b'\x89PNG\r\n\x1a\n')
        with patch('backend.app.previews.run_renderer', side_effect=png) as render:
            response = self.reviewer.get(url + '/redline/pages/3')
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response.headers['content-type'], 'image/png')
            self.assertIn('no-store', response.headers['cache-control'])
            self.assertEqual(self.reviewer.get(url + '/redline/pages/3').status_code, 200)
            self.assertEqual(render.call_count, 1)
        self.assertEqual(self.detail(p)['review_task']['status'], 'in_review')
        self.reviewer.post('/api/v1/auth/logout')
        self.assertEqual(self.reviewer.get(url + '/redline/pages/3').status_code, 401)

    def test_partial_render_failure_is_retryable(self):
        p = self.comparison(); url = f"/api/v1/comparisons/{p['comparison_id']}/preview"
        def fail_redline(source, destination):
            if source.name == 'redline.docx': raise RuntimeError('Failed renderer')
            return self.fake_pdf(source, destination)
        with patch('backend.app.previews.render_pdf', side_effect=fail_redline):
            self.assertEqual(self.reviewer.post(url).status_code, 503)
        self.assertEqual(self.reviewer.get(url + '/original/pages/1').status_code, 409)
        with patch('backend.app.previews.render_pdf', side_effect=self.fake_pdf):
            self.assertEqual(self.reviewer.post(url).status_code, 200)
        self.assertEqual(self.creator.get(p['versions'][0]['download_url']).content, b'original')


@unittest.skipUnless(shutil.which('libreoffice') and shutil.which('pdftoppm'), 'Native rendering is tested in Docker')
class NativePreviewTests(unittest.TestCase):
    def test_multipage_original_and_visual_redline_render(self):
        from backend.app.comparison.visual_redline import generate_visual_redline
        from backend.app.comparison.semantic_comparator import compare_semantic_changes
        with tempfile.TemporaryDirectory() as scratch:
            root = Path(scratch)
            original = root / 'original.docx'; revised = root / 'revised.docx'; redline = root / 'redline.docx'
            doc = Document(); doc.add_paragraph('Payment within 30 days.'); doc.add_page_break(); doc.add_paragraph('Second page.'); doc.save(original)
            doc = Document(); doc.add_paragraph('Payment within 14 days.'); doc.add_page_break(); doc.add_paragraph('Second page.'); doc.save(revised)
            # Native preview uses the same visual redline file as downloads.
            generate_visual_redline(original, revised, compare_semantic_changes(original, revised), redline)
            counts = prepare_preview(root / 'cache', {'original': original, 'redline': redline})
            self.assertEqual(counts['original'], 2)
            self.assertGreaterEqual(counts['redline'], 2)
            for kind in counts:
                for page in (1, counts[kind]):
                    image = render_page(root / 'cache', kind, page)
                    self.assertTrue(image.read_bytes().startswith(b'\x89PNG\r\n\x1a\n'))
